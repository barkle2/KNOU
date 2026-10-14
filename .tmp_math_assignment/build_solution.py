from pathlib import Path
import math
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT.parent / '[6-2]데이터과학을 위한 수학' / '중간대체과제_풀이.docx'
FONT = '맑은 고딕'

def el(tag, *children):
    node = OxmlElement('m:' + tag)
    for child in children:
        node.append(child)
    return node

def mr(text):
    node = el('r')
    props = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    for key in ['ascii', 'hAnsi', 'eastAsia']:
        fonts.set(qn('w:' + key), 'Cambria Math')
    props.append(fonts)
    size = OxmlElement('w:sz')
    size.set(qn('w:val'), '22')
    props.append(size)
    node.append(props)
    t = el('t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    node.append(t)
    return node

COMMANDS = {
    'ln': 'ln', 'exp': 'exp', 'partial': '∂', 'infty': '∞',
    'in': '∈', 'mathbb': '', 'le': '≤', 'ge': '≥', 'ne': '≠',
    'approx': '≈', 'cdot': '·', 'times': '×', 'mid': '|',
    'Rightarrow': '⇒', 'nabla': '∇', 'ell': 'ℓ', 'quad': '  ',
    'qquad': '    ', 'pm': '±', 'forall': '∀', 'sigma': 'σ', 'prime': '′',
}

class MathParser:
    """Small deterministic OMML authoring path for this document's notation."""
    def __init__(self, source):
        self.s = source
        self.i = 0

    def group(self):
        while self.i < len(self.s) and self.s[self.i] == ' ':
            self.i += 1
        if self.i < len(self.s) and self.s[self.i] == '{':
            self.i += 1
            return self.parse('}')
        return [self.atom()]

    def atom(self):
        c = self.s[self.i]
        self.i += 1
        if c == '{':
            return el('box', el('e', *self.parse('}')))
        if c != '\\':
            return mr('−' if c == '-' else c)
        match = re.match(r'[A-Za-z]+', self.s[self.i:])
        if not match:
            c = self.s[self.i]
            self.i += 1
            return mr(' ' if c in ',;' else c)
        cmd = match.group()
        self.i += len(cmd)
        if cmd == 'frac':
            return el('f', el('num', *self.group()), el('den', *self.group()))
        if cmd in ['left', 'right']:
            return self.atom()
        if cmd in ['mathrm', 'operatorname', 'mathbb']:
            return el('box', el('e', *self.group()))
        if cmd == 'int':
            return mr('∫')
        if cmd not in COMMANDS:
            raise ValueError('Unknown command: ' + cmd)
        return mr(COMMANDS[cmd])

    def parse(self, end=None):
        nodes = []
        while self.i < len(self.s):
            if self.s[self.i] == end:
                self.i += 1
                break
            atom = self.atom()
            sub = sup = None
            while self.i < len(self.s) and self.s[self.i] in '_^':
                kind = self.s[self.i]
                self.i += 1
                value = self.group()
                if kind == '_': sub = value
                else: sup = value
            if sub is not None and sup is not None:
                atom = el('sSubSup', el('e', atom), el('sub', *sub), el('sup', *sup))
            elif sub is not None:
                atom = el('sSub', el('e', atom), el('sub', *sub))
            elif sup is not None:
                atom = el('sSup', el('e', atom), el('sup', *sup))
            nodes.append(atom)
        return nodes

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2.0), Cm(2.0)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[name]
    st.font.name = FONT
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(6)
for name, size in [('Title', 21), ('Heading 1', 15), ('Heading 2', 12)]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.paragraph_format.space_before = Pt(10 if name != 'Title' else 0)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True

def p(text, bold=False):
    para = doc.add_paragraph()
    para.add_run(text).bold = bold
    return para

def eq(source):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after = Pt(7)
    para.paragraph_format.keep_together = True
    math_node = el('oMath', *MathParser(source).parse())
    para._p.append(math_node)
    return para

def h(text, level=2):
    doc.add_heading(text, level)

def page():
    doc.add_page_break()

def region_image(event=False):
    W,H=1100,620
    image=Image.new('RGB',(W,H),'white')
    d=ImageDraw.Draw(image)
    f=ImageFont.truetype('C:/Windows/Fonts/malgun.ttf',27)
    small=ImageFont.truetype('C:/Windows/Fonts/malgun.ttf',23)
    x0,y0,scale=140,535,108
    def xy(x,y):return (int(x0+x*scale),int(y0-y*scale))
    if event:
        d.polygon([xy(0,0),xy(4.2,4.2),xy(3.2,4.2),xy(0,1)],fill='#DAE6F0')
    else:
        d.polygon([xy(0,0),xy(4.2,4.2),xy(0,4.2)],fill='#DAE6F0')
    for n in range(1,5):
        d.line([xy(n,0),xy(n,4.3)], fill='#E1E1E1',width=1)
        d.line([xy(0,n),xy(4.4,n)], fill='#E1E1E1',width=1)
        d.text((xy(n,0)[0]-7,y0+9),str(n),font=small,fill='black')
        d.text((x0-34,xy(0,n)[1]-15),str(n),font=small,fill='black')
    d.line([xy(0,0),xy(4.55,0)],fill='black',width=3)
    d.line([xy(0,0),xy(0,4.55)],fill='black',width=3)
    d.polygon([xy(4.55,0),(xy(4.55,0)[0]-14,y0-6),(xy(4.55,0)[0]-14,y0+6)],fill='black')
    d.polygon([xy(0,4.55),(x0-6,xy(0,4.55)[1]+14),(x0+6,xy(0,4.55)[1]+14)],fill='black')
    # Dashed boundaries indicate that the positive-density region is open.
    def dashed(a,b):
        A=xy(*a);B=xy(*b)
        for k in range(0,80,2):
            d.line([(A[0]+(B[0]-A[0])*k/80,A[1]+(B[1]-A[1])*k/80),
                    (A[0]+(B[0]-A[0])*(k+1)/80,A[1]+(B[1]-A[1])*(k+1)/80)],fill='#304E65',width=3)
    dashed((0,0),(4.2,4.2))
    d.text((610,69),'y = x',font=f,fill='black')
    if event:
        dashed((0,1),(3.2,4.2))
        d.text((343,36),'y = x + 1',font=f,fill='black')
        d.text((700,230),'적분 영역 A',font=f,fill='black')
        d.text((700,275),'x > 0',font=f,fill='black')
        d.text((700,315),'x < y ≤ x + 1',font=f,fill='black')
    else:
        d.text((700,230),'양의 밀도 영역 D',font=f,fill='black')
        d.text((700,275),'0 < x < y',font=f,fill='black')
        d.text((700,315),'위쪽으로 무한히 이어짐',font=small,fill='black')
    d.text((635,y0-2),'x',font=f,fill='black')
    d.text((x0-30,23),'y',font=f,fill='black')
    d.text((x0-29,y0+7),'0',font=small,fill='black')
    path=ROOT/('event_region.png' if event else 'density_region.png')
    image.save(path)
    return path

def figure(event=False):
    para=doc.add_paragraph()
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(region_image(event)),width=Cm(14.0))
    para.paragraph_format.space_after=Pt(3)
    para=p('그림 2  대기시간이 1시간 이하인 영역' if event else '그림 1  결합밀도가 양수인 영역')
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.runs[0].font.size=Pt(9)

doc.add_paragraph('데이터과학을 위한 수학\n중간대체과제 풀이',style='Title')
p('로지스틱 손실함수의 편미분과 테일러 근사, 요청 발생시간의 결합밀도에 관한 두 문제를 문항 순서에 따라 풀이한다. 각 소문항에 계산 과정과 최종 결과를 제시한다.')
h('문제 1 로지스틱 손실함수와 테일러 근사',1)
p('두 학습자료는 (x₁, y₁) = (1, 1), (x₂, y₂) = (−1, −1)이다. 주어진 목적함수는 다음과 같다.')
eq(r'L(w,b)=\frac{1}{2}\{\ln(1+e^{-w-b})+\ln(1+e^{-w+b})\}+\frac{1}{4}(w^2+b^2)')
h('1번 소문항 편도함수와 원점에서의 값')
p('합성함수의 미분법을 적용한다. w에 대한 미분에서는 두 지수의 미분이 모두 −1이고, b에 대한 미분에서는 각각 −1과 1이다. 벌점항의 편도함수는 각각 w/2와 b/2이다.')
eq(r'\frac{\partial L}{\partial w}=-\frac{1}{2}\left(\frac{1}{1+e^{w+b}}+\frac{1}{1+e^{w-b}}\right)+\frac{w}{2}')
eq(r'\frac{\partial L}{\partial b}=\frac{1}{2}\left(-\frac{1}{1+e^{w+b}}+\frac{1}{1+e^{w-b}}\right)+\frac{b}{2}')
p('w = b = 0을 대입하면 각 분모가 2가 되므로 다음 값을 얻는다.')
eq(r'\frac{\partial L}{\partial w}(0,0)=-\frac{1}{2},\qquad \frac{\partial L}{\partial b}(0,0)=0')
h('2번 소문항 도함수와 2차 테일러 다항식')
eq(r'g(z)=\ln(1+e^z),\qquad g\prime(z)=\frac{e^z}{1+e^z}')
eq(r'g\prime\prime(z)=\frac{e^z(1+e^z)-e^{2z}}{(1+e^z)^2}=\frac{e^z}{(1+e^z)^2}')
eq(r'g(0)=\ln 2,\qquad g\prime(0)=\frac{1}{2},\qquad g\prime\prime(0)=\frac{1}{4}')
p('원점을 중심으로 하는 2차 테일러 다항식의 정의에 대입한다.')
eq(r'P_2(z)=g(0)+g\prime(0)z+\frac{g\prime\prime(0)}{2}z^2=\ln 2+\frac{z}{2}+\frac{z^2}{8}')

page()
h('문제 1 이차 근사와 최소점 판단',1)
h('3번 소문항 이차 근사함수')
p('P₂에 z = −w−b와 z = −w+b를 대입하고, 벌점항은 원래 형태로 유지한다.')
eq(r'Q(w,b)=\frac{1}{2}\{P_2(-w-b)+P_2(-w+b)\}+\frac{1}{4}(w^2+b^2)')
eq(r'Q(w,b)=\ln 2-\frac{w}{2}+\frac{(w+b)^2+(w-b)^2}{16}+\frac{w^2+b^2}{4}')
p('(w+b)²+(w−b)² = 2w²+2b²이므로 교차항이 소거된다.')
eq(r'Q(w,b)=\ln 2-\frac{w}{2}+\frac{3}{8}w^2+\frac{3}{8}b^2')
h('4번 소문항 근사함수의 정지점')
eq(r'\frac{\partial Q}{\partial w}=-\frac{1}{2}+\frac{3}{4}w,\qquad \frac{\partial Q}{\partial b}=\frac{3}{4}b')
p('두 편도함수를 동시에 0으로 놓으면 다음 점을 얻는다. w²와 b²의 계수가 모두 양수이므로 이 점은 Q의 유일한 전역 최소점이다.')
eq(r'(w_0,b_0)=\left(\frac{2}{3},0\right)')
h('5번 소문항 원래 손실함수의 최소점인지 판단')
p('1번 소문항에서 구한 원래 함수의 편도함수에 (2/3, 0)을 대입한다.')
eq(r'\frac{\partial L}{\partial w}\left(\frac{2}{3},0\right)=\frac{1}{3}-\frac{1}{1+e^{2/3}}\approx -0.0059103\ne 0')
eq(r'\frac{\partial L}{\partial b}\left(\frac{2}{3},0\right)=0')
p('따라서 (w₀, b₀)는 L의 최소점이 아니다. L은 실수 평면 전체에서 미분 가능하므로 최소점에서는 두 편도함수가 모두 0이어야 한다. 이 점에서 w 방향의 편도함수는 음수이므로, b를 0으로 유지하고 w를 조금 증가시키면 L이 감소한다.',bold=True)
p('테일러 다항식은 전개점 근처의 함수값과 도함수를 근사한다. 고차항을 생략한 Q의 최소점이 원래 L의 정지점과 일치한다는 보장은 없다. 추정값을 얻은 뒤 원래 함수의 기울기나 목적함수 값을 확인하고, 필요한 경우 원래 함수에 대한 수치 최적화로 보정해야 한다.')
p('참고로 L은 로지스틱 손실의 볼록성에 양의 이차 벌점이 더해져 엄격히 볼록하며, 벌점으로 인해 파라미터의 크기가 커질수록 무한히 증가한다. 따라서 유일한 전역 최소점이 존재한다. b에 대한 대칭성으로 최소점의 b는 0이고, w는 다음 방정식의 해이다.')
eq(r'\frac{w}{2}=\frac{1}{1+e^w},\qquad (w^*,b^*)\approx(0.6748316,0)')

page()
h('문제 2 요청 발생시간의 결합확률밀도',1)
p('X와 Y는 관측 시작 후 첫 번째와 두 번째 요청까지의 시간이다. 시간 단위는 시간이며, 주어진 결합밀도는 다음과 같다.')
eq(r'f(x,y)=Ce^{-y}\quad(0<x<y<\infty),\qquad C>0')
p('영역 밖에서는 f(x, y) = 0이다.')
h('1번 소문항 적분 영역과 정규화 상수')
p('양의 밀도 영역 D는 제1사분면에서 직선 y = x의 위쪽이다. x = 0과 y = x의 경계는 포함하지 않는다. 경계는 면적이 0이므로 적분값에는 영향을 주지 않는다.')
eq(r'D=\{(x,y):0<x<y<\infty\}')
figure()
p('x를 먼저 적분하면 y > 0을 고정했을 때 0 < x < y이다.')
eq(r'1=\int_0^{\infty}\int_0^y Ce^{-y}\,dx\,dy=C\int_0^{\infty}ye^{-y}\,dy')
p('부분적분에서 u = y, dv = e⁻ʸdy로 놓으면 v = −e⁻ʸ이다.')
eq(r'\int_0^{\infty}ye^{-y}\,dy=[-(y+1)e^{-y}]_0^{\infty}=1')
p('따라서 x를 먼저 적분한 결과는 C = 1이다.')
p('y를 먼저 적분하면 x > 0을 고정했을 때 x < y < ∞이다.')
eq(r'1=\int_0^{\infty}\int_x^{\infty}Ce^{-y}\,dy\,dx=C\int_0^{\infty}e^{-x}\,dx=C')
p('두 적분 순서 모두 정규화 상수 C = 1을 준다.',bold=True)

page()
h('문제 2 주변밀도와 대기시간의 확률',1)
h('2번 소문항 주변확률밀도와 독립 여부')
p('정규화 상수 C = 1을 사용한다. X의 주변밀도는 y에 대해 적분하고, Y의 주변밀도는 x에 대해 적분하여 구한다.')
eq(r'f_X(x)=\int_x^{\infty}e^{-y}\,dy=e^{-x}\quad(x>0)')
eq(r'f_Y(y)=\int_0^ye^{-y}\,dx=ye^{-y}\quad(y>0)')
p('x ≤ 0에서는 fX(x) = 0이고, y ≤ 0에서는 fY(y) = 0이다. X는 지수분포의 밀도를, Y는 형태모수 2와 비율모수 1인 감마분포의 밀도를 갖는다.')
eq(r'f_X(x)f_Y(y)=ye^{-(x+y)}\quad(x>0,\ y>0)')
p('이 곱은 결합밀도 e⁻ʸ와 일반적으로 같지 않다. 특히 0 < y < x인 영역에서는 결합밀도가 0인 반면 주변밀도의 곱은 양수이다. 따라서 X와 Y는 독립이 아니다.',bold=True)
p('문제 설명에서 서로 독립이라고 한 것은 연속하는 요청 사이의 대기시간이다. 관측 시작부터의 누적 발생시간 X와 Y는 항상 X < Y라는 제약을 공유한다.')
h('3번 소문항 1시간 이내에 두 번째 요청이 발생할 확률')
p('사건 Y−X ≤ 1과 원래 밀도의 영역을 동시에 만족하는 영역은 다음과 같다. 직선 y = x와 y = x+1 사이의 띠 모양 영역이다.')
eq(r'A=\{(x,y):x>0,\ x<y\le x+1\}')
figure(event=True)
p('y를 먼저 적분하는 반복적분은 다음과 같다.')
eq(r'P(Y-X\le1)=\int_0^{\infty}\int_x^{x+1}e^{-y}\,dy\,dx')
p('x를 먼저 적분할 때는 y = 1을 기준으로 나눈다. 0 < y ≤ 1에서는 0 < x < y이고, y > 1에서는 y−1 ≤ x < y이다.')
eq(r'P(Y-X\le1)=\int_0^1\int_0^ye^{-y}\,dx\,dy+\int_1^{\infty}\int_{y-1}^ye^{-y}\,dx\,dy')

page()
h('문제 2 확률 계산과 조건부 예측',1)
h('3번 소문항 확률 계산')
p('앞에서 세운 반복적분 중 y를 먼저 적분하는 방법으로 계산한다.')
eq(r'P(Y-X\le1)=\int_0^{\infty}[e^{-x}-e^{-(x+1)}]\,dx')
eq(r'=(1-e^{-1})\int_0^{\infty}e^{-x}\,dx=1-e^{-1}\approx0.6321206')
p('즉, 첫 번째 요청 이후 1시간 이내에 두 번째 요청이 발생할 확률은 약 63.21%이다.',bold=True)
h('4번 소문항 조건부확률밀도와 조건부 기대값')
p('x > 0을 고정하고 결합밀도를 X의 주변밀도로 나눈다.')
eq(r'f_{Y\mid X}(y\mid x)=\frac{f(x,y)}{f_X(x)}=\frac{e^{-y}}{e^{-x}}=e^{-(y-x)}\quad(y>x)')
p('y ≤ x에서는 조건부밀도가 0이다. x ≤ 0에서는 fX(x) = 0이므로 이 비율식으로 조건부밀도를 정의하지 않는다. x > 0에서 조건부밀도의 적분값은 1이다.')
eq(r'\int_x^{\infty}e^{-(y-x)}\,dy=1')
p('조건부 기대값의 정의에 따라 다음 적분을 계산한다.')
eq(r'm(x)=E[Y\mid X=x]=\int_x^{\infty}ye^{-(y-x)}\,dy=e^x\int_x^{\infty}ye^{-y}\,dy')
p('부분적분에서 u = y, dv = e⁻ʸdy, du = dy, v = −e⁻ʸ로 놓는다.')
eq(r'\int_x^{\infty}ye^{-y}\,dy=[-ye^{-y}]_x^{\infty}+\int_x^{\infty}e^{-y}\,dy')
eq(r'=xe^{-x}+e^{-x}=(x+1)e^{-x}')
p('따라서 조건부 기대값은 다음과 같다.')
eq(r'm(x)=e^x(x+1)e^{-x}=x+1\quad(x>0)')
p('첫 번째 요청이 관측 시작 2시간 후에 발생했다면 x = 2이다.')
eq(r'E[Y\mid X=2]=m(2)=3')
eq(r'E[Y-X\mid X=2]=m(2)-2=1')
p('두 번째 요청의 예측 발생시간은 관측 시작 3시간 후이고, 첫 번째 요청 이후의 예측 추가 대기시간은 1시간이다.',bold=True)

doc.core_properties.title='데이터과학을 위한 수학 중간대체과제 풀이'
doc.core_properties.subject='로지스틱 손실의 미분과 테일러 근사 및 결합확률밀도의 적분'
doc.core_properties.author=''
doc.save(OUTPUT)
print(OUTPUT)
print('paragraphs:',len(doc.paragraphs),'equations:',len(doc._element.xpath('//m:oMath')))
